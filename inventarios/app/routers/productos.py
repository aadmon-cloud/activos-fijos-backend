# app/routers/productos.py
from fastapi import APIRouter, Depends, HTTPException, Query, Request, UploadFile, File, Form
from typing import List, Optional
from datetime import datetime
from .. import crud
from ..models import ProductoCreate, ProductoUpdate
from ..database import get_database
from ..utils.codigos import generar_codigo_barras, generar_qr_code
from motor.motor_asyncio import AsyncIOMotorDatabase
import pandas as pd
import io
import chardet
import os
import shutil
from bson import ObjectId
from ..core.dependencies import require_authenticated, require_super_admin
from fastapi.responses import StreamingResponse
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

router = APIRouter(prefix="/productos", tags=["productos"])

def convertir_producto(producto):
    """Convertir ObjectId a string para JSON de forma segura"""
    if producto is None:
        return None
    
    # Crear una copia para no modificar el original
    producto_convertido = dict(producto)
    
    # Convertir _id a id si existe
    if '_id' in producto_convertido:
        producto_convertido['id'] = str(producto_convertido['_id'])
        del producto_convertido['_id']
    
    return producto_convertido

# ==========================================
# DESCARGA DE PLANTILLA EXCEL
# ==========================================

@router.get("/plantilla-excel")
async def descargar_plantilla_excel():
    """Descargar plantilla Excel para carga masiva de productos"""
    
    # Crear libro de trabajo
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Productos"
    
    # Definir encabezados
    headers = ["codigo", "nombre", "descripcion", "categoria", "stock_minimo", "precio_unitario", "ubicacion"]
    
    # Estilos para encabezados
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Escribir encabezados
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    # Agregar ejemplo de datos (fila 2)
    ejemplos = [
        "PROD001", "Producto de ejemplo", "Descripción del producto", 
        "Electrónica", "5", "100.00", "Almacén Principal"
    ]
    
    for col, ejemplo in enumerate(ejemplos, 1):
        cell = ws.cell(row=2, column=col, value=ejemplo)
        cell.border = thin_border
    
    # Agregar fila con notas (fila 3)
    nota = ws.cell(row=3, column=1, value="NOTA: No modificar los encabezados. El código debe ser único.")
    nota.font = Font(italic=True, size=10, color="FF0000")
    
    # Ajustar anchos de columnas
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 20
    
    # Guardar en memoria
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    # Retornar como archivo descargable
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=plantilla_productos.xlsx"}
    )
@router.post("/cargar-excel")
async def cargar_productos_excel(
    archivo: UploadFile = File(...),
    db: AsyncIOMotorDatabase = Depends(get_database)
):
    """Carga masiva de productos desde archivo Excel o CSV"""
    try:
        # Validar extensión
        if not (archivo.filename.endswith('.xlsx') or 
                archivo.filename.endswith('.xls') or 
                archivo.filename.endswith('.csv')):
            raise HTTPException(400, "Formato no soportado. Use .xlsx, .xls o .csv")
        
        print(f"📁 Procesando archivo: {archivo.filename}")
        
        # Leer archivo
        contents = await archivo.read()
        
        # Procesar según extensión
        if archivo.filename.endswith('.csv'):
            encoding = chardet.detect(contents)['encoding'] or 'utf-8'
            print(f"📄 Encoding detectado: {encoding}")
            df = pd.read_csv(io.BytesIO(contents), encoding=encoding)
        else:
            df = pd.read_excel(io.BytesIO(contents))
        
        print(f"📊 Filas leídas: {len(df)}")
        
        # Validar columnas obligatorias
        if 'codigo' not in df.columns:
            raise HTTPException(400, "El archivo debe tener una columna 'codigo'")
        if 'nombre' not in df.columns:
            raise HTTPException(400, "El archivo debe tener una columna 'nombre'")
        
        # Validar códigos duplicados en Excel
        codigos_excel = df['codigo'].astype(str).str.strip().tolist()
        codigos_duplicados = [c for c in set(codigos_excel) if codigos_excel.count(c) > 1]
        if codigos_duplicados:
            raise HTTPException(400, f"Códigos duplicados en el Excel: {', '.join(codigos_duplicados[:5])}")
        
        # Validar códigos existentes en BD
        codigos_existentes = []
        productos_col = db["productos"]
        
        for codigo in codigos_excel:
            if codigo and codigo != 'nan':
                existe = await productos_col.find_one({"codigo": codigo})
                if existe:
                    codigos_existentes.append(codigo)
        
        if codigos_existentes:
            raise HTTPException(400, f"Los siguientes códigos ya existen: {', '.join(codigos_existentes[:5])}")
        
        # Procesar productos
        resultados = []
        exitosos = 0
        fallidos = 0
        
        for idx, row in df.iterrows():
            try:
                codigo = str(row['codigo']).strip()
                if not codigo or codigo == 'nan':
                    raise ValueError(f"Código vacío en fila {idx+2}")
                
                nombre = str(row['nombre']).strip()
                if not nombre or nombre == 'nan':
                    raise ValueError(f"Nombre vacío en fila {idx+2}")
                
                producto_data = {
                    "codigo": codigo,
                    "nombre": nombre,
                    "descripcion": str(row.get('descripcion', '')) if pd.notna(row.get('descripcion', '')) else None,
                    "categoria": str(row.get('categoria', '')) if pd.notna(row.get('categoria', '')) else None,
                    "stock_minimo": int(row.get('stock_minimo', 0)) if pd.notna(row.get('stock_minimo', 0)) else 0,
                    "stock_actual": 0,
                    "precio_unitario": float(row.get('precio', 0)) if pd.notna(row.get('precio', 0)) else 0.0,
                    "ubicacion": str(row.get('ubicacion', '')) if pd.notna(row.get('ubicacion', '')) else None,
                    "fecha_creacion": datetime.now(),
                    "fecha_actualizacion": None
                }
                
                result = await productos_col.insert_one(producto_data)
                
                resultados.append({
                    "codigo": codigo,
                    "nombre": nombre,
                    "exitoso": True,
                    "error": None
                })
                exitosos += 1
                
            except Exception as e:
                print(f"❌ Error fila {idx+2}: {str(e)}")
                resultados.append({
                    "codigo": row.get('codigo', f'Fila {idx+2}'),
                    "nombre": row.get('nombre', ''),
                    "exitoso": False,
                    "error": str(e)
                })
                fallidos += 1
        
        return {
            "total": len(resultados),
            "exitosos": exitosos,
            "fallidos": fallidos,
            "productos": resultados
        }
        
    except HTTPException:
        raise
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(500, f"Error procesando archivo: {str(e)}")

@router.get("/")
async def leer_productos(
    skip: int = 0,
    limit: int = 100,
    db: AsyncIOMotorDatabase = Depends(get_database)
):
    """Obtener lista de todos los productos"""
    try:
        print("🔵 Obteniendo productos...")
        productos = await crud.get_productos(db, skip=skip, limit=limit)
        print(f"✅ Productos encontrados: {len(productos)}")
        
        return productos
    except Exception as e:
        print(f"❌ Error en leer_productos: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/buscar")
async def buscar_productos(
    q: str = Query(..., min_length=1, description="Término de búsqueda"),
    db: AsyncIOMotorDatabase = Depends(get_database)
):
    """Buscar productos por nombre, código o descripción"""
    try:
        print(f"🔵 Buscando: {q}")
        productos = await crud.buscar_productos(db, query=q)
        print(f"✅ Encontrados: {len(productos)}")
        
        # Convertir cada producto de forma segura
        productos_convertidos = []
        for producto in productos:
            producto_conv = convertir_producto(producto)
            if producto_conv:
                productos_convertidos.append(producto_conv)
        
        return productos_convertidos
    except Exception as e:
        print(f"❌ Error en buscar_productos: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/{producto_id}")
async def leer_producto(
    producto_id: str,
    db: AsyncIOMotorDatabase = Depends(get_database)
):
    """Obtener un producto por su ID o código"""
    try:
        db_producto = await crud.get_producto(db, producto_id=producto_id)
        if db_producto is None:
            raise HTTPException(status_code=404, detail="Producto no encontrado")
        
        return convertir_producto(db_producto)
    except Exception as e:
        print(f"❌ Error en leer_producto: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/codigo/{codigo}")
async def leer_producto_por_codigo(
    codigo: str,
    db: AsyncIOMotorDatabase = Depends(get_database)
):
    """Obtener un producto por su código"""
    try:
        db_producto = await crud.get_producto_por_codigo(db, codigo=codigo)
        if db_producto is None:
            raise HTTPException(status_code=404, detail="Producto no encontrado")
        
        return convertir_producto(db_producto)
    except Exception as e:
        print(f"❌ Error en leer_producto_por_codigo: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/", status_code=201)
async def crear_producto(
    producto: ProductoCreate,
    db: AsyncIOMotorDatabase = Depends(get_database)
):
    """Crear un nuevo producto"""
    try:
        nuevo_producto = await crud.crear_producto(db=db, producto=producto)
        return convertir_producto(nuevo_producto)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        print(f"❌ Error en crear_producto: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@router.put("/{producto_id}")
async def actualizar_producto(
    producto_id: str,
    producto: ProductoUpdate,
    db: AsyncIOMotorDatabase = Depends(get_database)
):
    """Actualizar un producto existente"""
    try:
        db_producto = await crud.actualizar_producto(db, producto_id=producto_id, producto=producto)
        if db_producto is None:
            raise HTTPException(status_code=404, detail="Producto no encontrado")
        
        return convertir_producto(db_producto)
    except Exception as e:
        print(f"❌ Error en actualizar_producto: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@router.delete("/{producto_id}")
async def eliminar_producto(
    producto_id: str,
    db: AsyncIOMotorDatabase = Depends(get_database)
):
    """Eliminar un producto"""
    try:
        success = await crud.eliminar_producto(db, producto_id=producto_id)
        if not success:
            raise HTTPException(status_code=404, detail="Producto no encontrado")
        return {"message": "Producto eliminado exitosamente"}
    except Exception as e:
        print(f"❌ Error en eliminar_producto: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/{producto_id}/codigo-barras")
async def obtener_codigo_barras(
    producto_id: str,
    db: AsyncIOMotorDatabase = Depends(get_database)
):
    """Obtener imagen de código de barras para un producto"""
    try:
        producto = await crud.get_producto(db, producto_id=producto_id)
        if not producto:
            raise HTTPException(status_code=404, detail="Producto no encontrado")
        
        barcode = generar_codigo_barras(producto['codigo'])
        if not barcode:
            raise HTTPException(status_code=500, detail="Error generando código de barras")
        
        return {"codigo_barras": barcode}
    except Exception as e:
        print(f"❌ Error en codigo_barras: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/{producto_id}/qr-code")
async def obtener_qr_code(
    producto_id: str,
    db: AsyncIOMotorDatabase = Depends(get_database)
):
    """Obtener código QR para un producto"""
    try:
        producto = await crud.get_producto(db, producto_id=producto_id)
        if not producto:
            raise HTTPException(status_code=404, detail="Producto no encontrado")
        
        qr = generar_qr_code(producto['codigo'], {"nombre": producto['nombre']})
        if not qr:
            raise HTTPException(status_code=500, detail="Error generando código QR")
        
        return {"qr_code": qr}
    except Exception as e:
        print(f"❌ Error en qr_code: {e}")
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/{identificador}/con-ubicaciones")
async def obtener_producto_con_ubicaciones(
    identificador: str,
    session_id: str = None,
    db: AsyncIOMotorDatabase = Depends(get_database)
):
    try:
        print(f"🔵 Buscando producto: {identificador}")
        
        producto = await crud.get_producto_por_codigo(db, identificador)
        if not producto:
            from bson import ObjectId
            if ObjectId.is_valid(identificador):
                producto = await crud.get_producto(db, identificador)
        
        if not producto:
            return {"error": "Producto no encontrado"}
        
        print(f"✅ Producto encontrado: {producto.get('codigo')}")
        
        # Buscar ubicaciones
        ubicaciones_col = db["ubicaciones"]
        codigo_busqueda = producto["codigo"]
        print(f"🔍 Buscando ubicaciones con producto_codigo: {codigo_busqueda}")
        
        ubicaciones = await ubicaciones_col.find({"producto_codigo": codigo_busqueda}).to_list(length=None)
        print(f"📍 Ubicaciones encontradas: {len(ubicaciones)}")
        
        for ub in ubicaciones:
            print(f"   - {ub.get('ubicacion')}: stock_fisico={ub.get('stock_fisico')}")
        
        # Obtener reserva del usuario actual
        reservas_col = db["reservas"]
        reserva_actual = None
        if session_id:
            reserva_actual = await reservas_col.find_one({"session_id": session_id})
        
        # Calcular stock_reservado_actual por ubicación
        stock_reservado_actual_por_ubicacion = {}
        if reserva_actual:
            for item in reserva_actual.get('items', []):
                ubicacion = item.get('ubicacion')
                if ubicacion:
                    stock_reservado_actual_por_ubicacion[ubicacion] = stock_reservado_actual_por_ubicacion.get(ubicacion, 0) + item.get('cantidad', 0)
        
        # Procesar ubicaciones
        stock_por_ubicacion = []
        for ub in ubicaciones:
            ubicacion_nombre = ub.get("ubicacion")
            stock_fisico = ub.get("stock_fisico", 0)
            stock_reservado_ubicacion = ub.get("stock_reservado", 0)
            stock_reservado_actual = stock_reservado_actual_por_ubicacion.get(ubicacion_nombre, 0)
            
            # ✅ Calcular reservas de OTROS en esta ubicación
            otras_reservas = await reservas_col.find({
                "session_id": {"$ne": session_id} if session_id else {},
                "items": {
                    "$elemMatch": {
                        "producto_id": producto["codigo"],
                        "ubicacion": ubicacion_nombre
                    }
                }
            }).to_list(length=None)
            
            stock_reservado_otros = 0
            for reserva in otras_reservas:
                for item in reserva.get('items', []):
                    if item.get('producto_id') == producto["codigo"] and item.get('ubicacion') == ubicacion_nombre:
                        stock_reservado_otros += item.get('cantidad', 0)
            
            # Stock disponible
            stock_disponible = stock_fisico - stock_reservado_otros - stock_reservado_actual
            
            print(f"📍 {ubicacion_nombre}: físico={stock_fisico}, reservado_actual={stock_reservado_actual}, reservado_otros={stock_reservado_otros}, disponible={stock_disponible}")
            
            stock_por_ubicacion.append({
                "id": str(ub["_id"]),
                "ubicacion": ubicacion_nombre,
                "stock_fisico": stock_fisico,
                "stock_reservado": stock_reservado_ubicacion,
                "stock_reservado_actual": stock_reservado_actual,
                "stock_reservado_otros": stock_reservado_otros,
                "stock_disponible": max(0, stock_disponible)
            })
        
        return {
            "producto": producto,
            "ubicaciones": stock_por_ubicacion,
            "stock_por_ubicacion": stock_por_ubicacion,
            "total_ubicaciones": len(stock_por_ubicacion)
        }
        
    except HTTPException:
        raise
    except Exception as e:
        print(f"❌ Error en obtener_producto_con_ubicaciones: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Error interno: {str(e)}")

@router.post("/{producto_id}/ubicaciones")
async def crear_ubicacion_producto(
    producto_id: str,
    ubicacion: str = Query(..., min_length=1),
    cantidad: int = Query(0, ge=0),
    db: AsyncIOMotorDatabase = Depends(get_database),
    #current_user: dict = Depends(require_authenticated)
):
    """Crear una nueva ubicación para un producto"""
    try:
        print(f"🔵 Creando ubicación: producto={producto_id}, ubicacion={ubicacion}, cantidad={cantidad}")
        
        # Buscar producto
        producto = await crud.get_producto_por_codigo(db, producto_id)
        if not producto:
            try:
                from bson import ObjectId
                if ObjectId.is_valid(producto_id):
                    producto = await crud.get_producto(db, producto_id)
            except:
                pass
        
        if not producto:
            raise HTTPException(status_code=404, detail="Producto no encontrado")
        
        codigo_producto = producto["codigo"]

        usuario_actual = "admin"
        
        # Crear ubicación
        resultado = await crud.crear_ubicacion(db, codigo_producto, ubicacion, cantidad, usuario_actual)
        
        # ✅ Asegurar que el resultado sea serializable
        return {
            "success": True,
            "ubicacion": resultado,
            "message": f"Ubicación '{ubicacion}' creada con {cantidad} unidades"
        }
        
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        print(f"❌ Error: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))
    
@router.put("/{producto_id}/ubicaciones/{ubicacion}")
async def actualizar_ubicacion_producto(
    producto_id: str,
    ubicacion: str,
    cantidad: int = Query(..., ge=0),
    db: AsyncIOMotorDatabase = Depends(get_database)
):
    """Actualizar stock de una ubicación"""
    try:
        # Buscar producto por código o ID
        producto = await crud.get_producto_por_codigo(db, producto_id)
        if not producto:
            try:
                from bson import ObjectId
                if ObjectId.is_valid(producto_id):
                    producto = await crud.get_producto(db, producto_id)
            except:
                pass
        
        if not producto:
            raise HTTPException(status_code=404, detail="Producto no encontrado")
        
        codigo_producto = producto["codigo"]
        
        # Obtener ubicación actual
        ubicaciones_col = db["ubicaciones"]
        ubicacion_actual = await ubicaciones_col.find_one({
            "producto_codigo": codigo_producto,
            "ubicacion": ubicacion
        })
        
        if not ubicacion_actual:
            raise HTTPException(status_code=404, detail=f"Ubicación '{ubicacion}' no encontrada")
        
        cantidad_anterior = ubicacion_actual.get("stock_fisico", 0)
        diferencia = cantidad - cantidad_anterior
        
        # Actualizar ubicación
        await ubicaciones_col.update_one(
            {"_id": ubicacion_actual["_id"]},
            {"$set": {"stock_fisico": cantidad}}
        )
        
        # Actualizar stock total del producto
        productos_col = db["productos"]
        nuevo_stock_total = max(0, (producto.get("stock_actual", 0) or 0) + diferencia)
        await productos_col.update_one(
            {"codigo": codigo_producto},
            {"$set": {"stock_actual": nuevo_stock_total}}
        )
        
        return {"success": True, "diferencia": diferencia}
        
    except HTTPException:
        raise
    except Exception as e:
        print(f"❌ Error: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@router.delete("/{producto_id}/ubicaciones/{ubicacion}")
async def eliminar_ubicacion_producto(
    producto_id: str,
    ubicacion: str,
    db: AsyncIOMotorDatabase = Depends(get_database)
):
    """Eliminar una ubicación de un producto"""
    try:
        # Buscar producto por código o ID
        producto = await crud.get_producto_por_codigo(db, producto_id)
        if not producto:
            try:
                from bson import ObjectId
                if ObjectId.is_valid(producto_id):
                    producto = await crud.get_producto(db, producto_id)
            except:
                pass
        
        if not producto:
            raise HTTPException(status_code=404, detail="Producto no encontrado")
        
        codigo_producto = producto["codigo"]
        
        ubicaciones_col = db["ubicaciones"]
        ubicacion_actual = await ubicaciones_col.find_one({
            "producto_codigo": codigo_producto,
            "ubicacion": ubicacion
        })
        
        if not ubicacion_actual:
            raise HTTPException(status_code=404, detail=f"Ubicación '{ubicacion}' no encontrada")
        
        cantidad_a_eliminar = ubicacion_actual.get("stock_fisico", 0)
        
        # Eliminar ubicación
        await ubicaciones_col.delete_one({"_id": ubicacion_actual["_id"]})
        
        # Actualizar stock total del producto
        productos_col = db["productos"]
        nuevo_stock_total = max(0, (producto.get("stock_actual", 0) or 0) - cantidad_a_eliminar)
        await productos_col.update_one(
            {"codigo": codigo_producto},
            {"$set": {"stock_actual": nuevo_stock_total}}
        )
        
        return {"success": True}
        
    except HTTPException:
        raise
    except Exception as e:
        print(f"❌ Error: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/{producto_id}/transferir")
async def transferir_stock(
    producto_id: str,
    origen: str = Query(...),
    destino: str = Query(...),
    cantidad: int = Query(..., ge=1),
    db: AsyncIOMotorDatabase = Depends(get_database),
    #current_user: dict = Depends(require_authenticated)
):
    """Transferir stock entre ubicaciones del mismo producto"""
    try:
        producto = await crud.get_producto(db, producto_id)
        if not producto:
            raise HTTPException(status_code=404, detail="Producto no encontrado")
        
        codigo_producto = producto["codigo"]
        
        # ✅ NO usar current_user - usar valor fijo
        usuario_actual = "admin"
        
        # Llamar a la función de transferencia
        resultado = await crud.transferir_stock_entre_ubicaciones(
            db, codigo_producto, origen, destino, cantidad, usuario_actual
        )
        
        return {"success": True, "message": f"Transferidas {cantidad} unidades de {origen} a {destino}"}
        
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        print(f"❌ Error en transferir_stock: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/{codigo}/stock-por-ubicacion")
async def obtener_stock_por_ubicacion(
    codigo: str,
    session_id: str = None,
    db: AsyncIOMotorDatabase = Depends(get_database)
):
    """
    Obtener stock disponible por ubicación para un producto
    Considera reservas de otros usuarios Y del usuario actual
    """
    try:
        # Buscar producto por código
        producto = await crud.get_producto_por_codigo(db, codigo)
        if not producto:
            raise HTTPException(status_code=404, detail="Producto no encontrado")
        
        # Obtener ubicaciones del producto
        ubicaciones_col = db["ubicaciones"]
        ubicaciones = await ubicaciones_col.find({"producto_codigo": codigo}).to_list(length=None)
        
        # Obtener reservas
        reservas_col = db["reservas"]
        
        # Obtener reserva del usuario actual
        reserva_actual = None
        if session_id:
            reserva_actual = await reservas_col.find_one({"session_id": session_id})
        
        # Calcular stock por ubicación
        stock_por_ubicacion = []
        stock_total_disponible = 0
        
        for ub in ubicaciones:
            ubicacion_nombre = ub.get("ubicacion")
            stock_fisico = ub.get("stock_fisico", 0)
            
            # ✅ Calcular reservas en esta ubicación por OTROS usuarios
            otras_reservas = await reservas_col.find({
                "session_id": {"$ne": session_id} if session_id else {},
                "items": {
                    "$elemMatch": {
                        "producto_id": codigo,
                        "ubicacion": ubicacion_nombre
                    }
                }
            }).to_list(length=None)
            
            stock_reservado_otros = 0
            for reserva in otras_reservas:
                for item in reserva.get('items', []):
                    if item.get('producto_id') == codigo and item.get('ubicacion') == ubicacion_nombre:
                        stock_reservado_otros += item.get('cantidad', 0)
            
            # ✅ Calcular reservas del USUARIO ACTUAL en esta ubicación
            stock_reservado_actual = 0
            if reserva_actual:
                for item in reserva_actual.get('items', []):
                    if item.get('producto_id') == codigo and item.get('ubicacion') == ubicacion_nombre:
                        stock_reservado_actual += item.get('cantidad', 0)
            
            # Stock disponible = físico - reservas de otros - reservas del usuario actual
            stock_disponible = stock_fisico - stock_reservado_otros - stock_reservado_actual
            stock_total_disponible += max(0, stock_disponible)
            
            print(f"📍 {ubicacion_nombre}: físico={stock_fisico}, reservado_otros={stock_reservado_otros}, reservado_actual={stock_reservado_actual}, disponible={stock_disponible}")
            
            stock_por_ubicacion.append({
                "ubicacion": ubicacion_nombre,
                "stock_fisico": stock_fisico,
                "stock_reservado_otros": stock_reservado_otros,
                "stock_reservado_actual": stock_reservado_actual,
                "stock_disponible": max(0, stock_disponible)
            })
        
        return {
            "producto": {
                "codigo": producto["codigo"],
                "nombre": producto["nombre"],
                "stock_total": producto.get("stock_actual", 0)
            },
            "stock_por_ubicacion": stock_por_ubicacion,
            "stock_total_disponible": stock_total_disponible,
            "tiene_stock": stock_total_disponible > 0
        }
        
    except HTTPException:
        raise
    except Exception as e:
        print(f"❌ Error en stock-por-ubicacion: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))

# Configuración de carpetas
UPLOAD_DIR = "app/static/imagenes"
os.makedirs(UPLOAD_DIR, exist_ok=True)

@router.post("/{producto_id}/subir-imagen")
async def subir_imagen_producto(
    producto_id: str,
    archivo: UploadFile = File(...),  # ← Debe ser exactamente "archivo"
    db: AsyncIOMotorDatabase = Depends(get_database),
    # current_user: dict = Depends(require_authenticated)  # Temporalmente comentado
):
    try:
        print(f"🔵 Recibiendo imagen para producto: {producto_id}")
        print(f"📁 Archivo: {archivo.filename}, Tipo: {archivo.content_type}")
        
        # Obtener producto
        producto = await crud.get_producto(db, producto_id)
        if not producto:
            raise HTTPException(status_code=404, detail="Producto no encontrado")
        
        # Validar que sea categoría Souvenir
        if producto.get('categoria') != 'SOUVENIR':
            raise HTTPException(status_code=400, detail="Solo productos de la categoría 'SOUVENIR' pueden tener imagen")
        
        # Validar tipo de archivo
        if not archivo.content_type or not archivo.content_type.startswith('image/'):
            raise HTTPException(status_code=400, detail="Solo se permiten archivos de imagen")
        
        # Crear directorio si no existe
        import os
        UPLOAD_DIR = "app/static/imagenes"
        os.makedirs(UPLOAD_DIR, exist_ok=True)
        
        # Generar nombre único
        extension = archivo.filename.split('.')[-1]
        nombre_archivo = f"{producto_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.{extension}"
        ruta_archivo = os.path.join(UPLOAD_DIR, nombre_archivo)
        
        # Guardar archivo
        import shutil
        with open(ruta_archivo, "wb") as buffer:
            shutil.copyfileobj(archivo.file, buffer)
        
        # Guardar URL en MongoDB
        from bson import ObjectId
        imagen_url = f"/static/imagenes/{nombre_archivo}"
        await db["productos"].update_one(
            {"_id": ObjectId(producto_id)},
            {"$set": {"imagen_url": imagen_url, "fecha_actualizacion": datetime.now()}}
        )
        
        print(f"✅ Imagen guardada: {imagen_url}")
        
        return {"success": True, "imagen_url": imagen_url}
        
    except HTTPException:
        raise
    except Exception as e:
        print(f"❌ Error: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))


@router.delete("/{producto_id}/eliminar-imagen")
async def eliminar_imagen_producto(
    producto_id: str,
    db: AsyncIOMotorDatabase = Depends(get_database),
    current_user: dict = Depends(require_authenticated)
):
    """Eliminar imagen de un producto"""
    try:
        producto = await crud.get_producto(db, producto_id)
        if not producto:
            raise HTTPException(status_code=404, detail="Producto no encontrado")
        
        imagen_url = producto.get('imagen_url')
        if imagen_url:
            # Eliminar archivo físico
            ruta_archivo = imagen_url.replace('/static/imagenes/', '')
            ruta_completa = os.path.join(UPLOAD_DIR, ruta_archivo)
            if os.path.exists(ruta_completa):
                os.remove(ruta_completa)
        
        await db["productos"].update_one(
            {"_id": ObjectId(producto_id)},
            {"$set": {"imagen_url": None, "fecha_actualizacion": datetime.now()}}
        )
        
        return {"success": True}
        
    except Exception as e:
        print(f"❌ Error: {e}")
        raise HTTPException(status_code=500, detail=str(e))

