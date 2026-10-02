import io
from datetime import UTC, datetime

import openpyxl

from modules.movimientos.domain.entities.movimiento import Movimiento
from modules.movimientos.infrastructure.excel.movimientos_excel_exporter import generar_excel
from shared.domain.entities.audit_persona import AuditPersona


def _audit_persona():
    return AuditPersona(id_persona="admin-id", nombre="admin", fecha=datetime.now(UTC))


def _leer_columna(contenido: bytes, nombre_columna: str, fila: int = 2):
    wb = openpyxl.load_workbook(io.BytesIO(contenido))
    hoja = wb["Movimientos"]
    encabezados = [c.value for c in hoja[1]]
    return hoja.cell(row=fila, column=encabezados.index(nombre_columna) + 1).value


def test_sin_movimientos_genera_hoja_con_mensaje():
    contenido = generar_excel([])

    wb = openpyxl.load_workbook(io.BytesIO(contenido))
    hoja = wb["Movimientos"]
    assert hoja.cell(row=2, column=1).value == "No hay movimientos para exportar"


def test_etiqueta_entrada_y_salida():
    entrada = Movimiento.registrar_entrada(
        producto_codigo="PROD-001",
        producto_nombre="Arroz",
        cantidad=10,
        motivo=None,
        tipo_origen=None,
        origen_nombre=None,
        ubicacion=None,
        notas=None,
        creado_por=_audit_persona(),
    )
    salida = Movimiento.registrar_salida(
        producto_codigo="PROD-001",
        producto_nombre="Arroz",
        cantidad=5,
        motivo=None,
        cliente_destino=None,
        ubicacion="A10-1",
        notas=None,
        remision="RM-001",
        creado_por=_audit_persona(),
    )

    contenido = generar_excel([entrada, salida])

    assert _leer_columna(contenido, "Tipo", fila=2) == "ENTRADA"
    assert _leer_columna(contenido, "Tipo", fila=3) == "SALIDA"


def test_etiqueta_bitacora_de_ubicaciones_no_se_confunde_con_salida():
    creacion = Movimiento.registrar_creacion_ubicacion(
        producto_codigo="PROD-001",
        producto_nombre="Arroz",
        ubicacion="A10-1",
        cantidad=50,
        creado_por=_audit_persona(),
    )
    transferencia = Movimiento.registrar_transferencia_ubicacion(
        producto_codigo="PROD-001",
        producto_nombre="Arroz",
        ubicacion_origen="A10-1",
        ubicacion_destino="A10-2",
        cantidad=20,
        creado_por=_audit_persona(),
    )

    contenido = generar_excel([creacion, transferencia])

    assert _leer_columna(contenido, "Tipo", fila=2) == "CREACIÓN DE UBICACIÓN"
    assert _leer_columna(contenido, "Tipo", fila=3) == "TRANSFERENCIA ENTRE UBICACIONES"
